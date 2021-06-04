import os
import sys
import urllib.request
import json,re

# 네이버 search api
def searchNAPI(category,keyword,display,start,sort):
    client_id = "rKPa0Ubq8anIE1p_IDAg"
    client_secret = "AuyDMRd17F"
    # client_id = "5riZ8F3Wvyp5pqmUvZ86"
    # client_secret = "bVCTSk556w"
    # client_id = "y1aJ70K7H7QGZqo0N851"
    # client_secret = "upfgG6lYj8"
    encText = urllib.parse.quote(keyword)
    url = "https://openapi.naver.com/v1/search/"+category+".json?query="+encText+"&display="+display+"&start="+start+"&sort="+sort # json 결과
    try:
        request = urllib.request.Request(url)
        request.add_header("X-Naver-Client-Id",client_id)
        request.add_header("X-Naver-Client-Secret",client_secret)
        response = urllib.request.urlopen(request)
        rescode = response.getcode()

        if(rescode==200):
            response_body = response.read()
            return json.loads(response_body.decode('utf-8'))
        else:
            print("Error Code:" + rescode)
    except Exception as e:
        print("Error Code:",e)

    return False

def searchNAPI2(category,keyword,display,start,sort):
    client_id = "rKPa0Ubq8anIE1p_IDAg"
    client_secret = "AuyDMRd17F"
    # client_id = "5riZ8F3Wvyp5pqmUvZ86"
    # client_secret = "bVCTSk556w"
    encText = urllib.parse.quote(keyword)
    url = "https://openapi.naver.com/v1/search/"+category+".json?query="+encText+"&display="+display+"&start="+start+"&sort="+sort # json 결과
    try:
        request = urllib.request.Request(url)
        request.add_header("X-Naver-Client-Id",client_id)
        request.add_header("X-Naver-Client-Secret",client_secret)
        response = urllib.request.urlopen(request)
        rescode = response.getcode()

        if(rescode==200):
            response_body = response.read()
            return json.loads(response_body.decode('utf-8'))
        else:
            print("Error Code:" + rescode)
    except Exception as e:
        print("Error Code:",e)

    return False

# 네이버 api - url 줄임
def shortURL(shortUrl):
    # client_id = "rKPa0Ubq8anIE1p_IDAg"
    # client_secret = "AuyDMRd17F"
    client_id = "5riZ8F3Wvyp5pqmUvZ86"
    client_secret = "bVCTSk556w"
    encText = urllib.parse.quote(shortUrl)
    data = "url=" + encText
    url = "https://openapi.naver.com/v1/util/shorturl"
    request = urllib.request.Request(url)
    request.add_header("X-Naver-Client-Id",client_id)
    request.add_header("X-Naver-Client-Secret",client_secret)
    response = urllib.request.urlopen(request, data=data.encode("utf-8"))
    rescode = response.getcode()
    if(rescode==200):
        response_body = response.read()
        rUrl = json.loads(response_body.decode('utf-8'))['result']
        return rUrl['url']
    else:
        print("Error Code:" + rescode)
    return shortUrl

# 네이버 api - url 줄임
def shortURL2(shortUrl):
    client_id = "rKPa0Ubq8anIE1p_IDAg"
    client_secret = "AuyDMRd17F"
    # client_id = "5riZ8F3Wvyp5pqmUvZ86"
    # client_secret = "bVCTSk556w"
    encText = urllib.parse.quote(shortUrl)
    data = "url=" + encText
    url = "https://openapi.naver.com/v1/util/shorturl"
    request = urllib.request.Request(url)
    request.add_header("X-Naver-Client-Id",client_id)
    request.add_header("X-Naver-Client-Secret",client_secret)
    response = urllib.request.urlopen(request, data=data.encode("utf-8"))
    rescode = response.getcode()
    if(rescode==200):
        response_body = response.read()
        rUrl = json.loads(response_body.decode('utf-8'))['result']
        return rUrl['url']
    else:
        print("Error Code:" + rescode)
    return shortUrl

# 다음 search api
def searchDAPI(category,keyword,page):
    apikey = "KakaoAK 5ffc76bc940c42ccf03bf109f49c793f"
    encText = urllib.parse.quote(keyword)
    url = "https://dapi.kakao.com/v2/search/"+category+"?query="+encText+"&page="+page+"&size=50&sort=recency" # json 결과

    request = urllib.request.Request(url)
    request.add_header("Authorization",apikey)
    response = urllib.request.urlopen(request)
    rescode = response.getcode()

    if(rescode==200):
        response_body = response.read()
        json.loads(response_body.decode('utf-8'))
        return json.loads(response_body.decode('utf-8'))
    else:
        print("Error Code:" + rescode)

    return False

# DB 저장하는 함수
def insert(conn,*args):
    import pymysql
    import datetime
    result = False
    try:
        curs = conn.cursor(pymysql.cursors.DictCursor)
        now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        tableName = 'portal_data'
        data = {
            'portal_type': args[0],
            'portal_name': args[1],
            'portal_title': args[2],
            'deviceType': 1,
            'writer': args[3],
            'writeDate': args[4],
            'title_key': args[5],
            'keyword': args[6],
            'keyword_type': args[7],
            'url': args[8],
            'createDate': now,
            'updateDate':now
        }
        if data['title_key'] == '':
            del data['title_key']

        if args[0] == 'media' or args[0] == 'reporter':
            if args[0] == 'reporter':
                data['news_type'] = data['keyword_type']
                data['reporter_name'] = args[9]
                data['ME_rank'] = args[10]
                data['uid'] = args[11]
                data['media_subname'] = args[12]
                data['media_main'] = args[13]
                del data['keyword_type']
            if args[0] == 'media':
                data['news_type'] = data['keyword_type']
                data['reporter_name'] = args[9]
                data['ME_rank'] = args[10]
                data['uid'] = args[11]
                data['media_subname'] = args[12]
                data['media_main'] = args[13]
                del data['keyword_type']
            data['media_name'] = data['portal_name']
            data['media_title'] = data['portal_title']
            data['media_content'] = args[3]
            tableName = 'media_data'
            del(data['portal_type'],data['deviceType'],data['portal_name'],data['portal_title'],data['writer'])

        placeholders = ', '.join(['%s'] * len(data))
        columns = ', '.join(data.keys())
        sql = "INSERT INTO "+tableName+" ( %s ) VALUES ( %s );" % (columns, placeholders)
        curs.execute(sql, list(data.values()))
        conn.commit()
    except Exception as e:
        if e.args[0] != 1062:
            print("===========에러==========\n에러 : ",e,"\n",args,"\n===========에러==========")
        else:
            result = True
            conn.rollback()
    finally:
        return result

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
# 이름 가져오기
def getWriter(dic,pName,pType):
    driver = webdriver.Chrome("c:\python36\driver\chromedriver")
    try:
        for item in dic:
            # alert창 체크
            try:
                driver.get(item['url'])
                WebDriverWait(driver, 2).until(EC.alert_is_present())
                print('alert창이 있어 건너뜁니다.')
                driver.quit()
                driver = webdriver.Chrome("c:\python36\driver\chromedriver")
                continue
            except:
                pass

            uWriter = None
            if pType == 'cafe':
                try:
                    if pName == 'naver':
                        element = WebDriverWait(driver, 5).until(EC.presence_of_element_located((By.ID, "cafe_main")))
                        driver.switch_to.frame(element)
                        uWriter = driver.find_element(by='class name', value="p-nick").text
                    else:
                        element = WebDriverWait(driver, 5).until(EC.presence_of_element_located((By.ID, "down")))
                        driver.switch_to.frame(element)
                        uWriter = driver.find_element(by='class name', value="article_writer").find_element(by='class name', value="txt_point").text
                except Exception as e:
                    if pName == 'naver':
                        driver.switch_to_window("naver_login")
                        driver.close()
                        driver.switch_to_window("")
                    continue
            else:
                try:
                    if pName == 'naver':
                        element = WebDriverWait(driver, 5).until(EC.presence_of_element_located((By.ID, "mainFrame")))
                        driver.switch_to.frame(element)
                        uWriter = driver.find_element(By.CLASS_NAME, "bg-body").find_element(By.CLASS_NAME, "nick").text
                    else:
                        if url.find('daum') != -1:
                            element = WebDriverWait(driver, 5).until(EC.presence_of_element_located((By.XPATH, "//*[@id='ProfileMenuMain']/div/div[1]/a")))
                            uWriter = element.text
                except Exception as e:
                    continue

            if uWriter:
                import pymysql
                conn = pymysql.connect(host='49.247.5.169',user='overware',password='uni1004!@',db='union',charset='utf8')
                curs = conn.cursor(pymysql.cursors.DictCursor)
                try:
                    sql = "update portal_data set writer=%s where portal_type=%s and portal_name=%s and keyword=%s and url=%s;"
                    curs.execute(sql,(uWriter.strip(),pType,pName,item['keyword'],item['url']))
                    conn.commit()
                finally:
                    conn.close()
    except Exception as e:
        print('driverError:',e)
    finally:
        driver.quit()

#이모티콘 삭제 처리 1
def remove_emoji(data):
    emoji_pattern = re.compile("["
        u"\U0001F600-\U0001F64F"  # emoticons
        u"\U0001F300-\U0001F5FF"  # symbols & pictographs
        u"\U0001F680-\U0001F6FF"  # transport & map symbols
        u"\U0001F1E0-\U0001F1FF"  # flags (iOS)
                           "]+", flags=re.UNICODE)

    return emoji_pattern.sub(r'', data)

#이모티콘 삭제 처리 2
def remove_emoji2(data):
    re_pattern = re.compile(u'[^\u0000-\uFFFF]', re.UNICODE);
    return re_pattern.sub('', data)

# web url 확인
from openpyxl import load_workbook
def checkLink(url):
    delURL = []
    wb = load_workbook('../제외url.xlsx', read_only=True)
    ws = wb.get_sheet_by_name("Sheet2")
    for r in ws.rows:
        delURL.append(r[0].value)

    if any(url.find(s) != -1 for s in delURL):
        return False
    else:
        if url.find("viki") != -1:
            url = url.split("?")[0]
        elif url.find("heraldcorp") != -1:
            url = url.split("&RURL=")[0]
        # elif url.find(".club/")

    return url

# db에 널을 text setting
def setText(s,t):
    s = s.replace("&lt;", "<")
    s = s.replace("&gt;", ">")
    s = s.replace("&quot;", '"')
    s = s.replace("&apos;", "'")
    s = s.replace("&amp;", "&")
    s = s.replace("<b>","")
    s = s.replace("</b>","")
    s = s.replace("\r","")

    if t == 0:
        s = (len(s) > 49) and s[:47]+"…" or s
        s = remove_emoji(s)

    return s

# db에 넣을 keyword setting
def getPutKeyword(title,text,arr):
    key = {
        'find':0,
        'keyword':''
    }
    for item in arr:
        findnum = title.find(item.replace("영화",""))
        if key['keyword'] == '' and int(findnum) != -1:
            key.update({'find':findnum,'keyword':item})
        if key['keyword'] != '' and key['find'] >= int(findnum) and int(findnum) != -1:
            key.update({'find':findnum,'keyword':item})

    if key['keyword'] == '' and text is not None:
        for item in arr:
            findnum = text.find(item.replace("영화",""))
            if key['keyword'] == '' and int(findnum) != -1:
                key.update({'find':findnum,'keyword':item})
            if key['keyword'] != '' and key['find'] >= int(findnum) and int(findnum) != -1:
                key.update({'find':findnum,'keyword':item})

    return key['keyword']

# db에 넣을 keyword_type get
def getPutKeywordType(keyword,conn,curs):
    returnVal = None
    with conn.cursor() as curs:
        sql = 'SELECT keyword_type FROM keyword_data WHERE keyword = %s;'
        curs.execute(sql, (keyword))
        kType = curs.fetchone()
        if kType:
            returnVal = kType[0]
    return returnVal

# 내용 키워드 체크
def checkKeyword(title,text,add,delete,keyword=None):
    textResult = None
    titleResult = None

    if text:
        if any(text.find(s.replace("영화","")) != -1 for s in add):
            if all(text.find(s) == -1 for s in delete):
                textResult = True
            else:
                textResult = False
        else:
            textResult = False
    if title:
        if any(title.find(s.replace("영화","")) != -1 for s in add):
            if all(title.find(s) == -1 for s in delete):
                titleResult = True
            else:
                titleResult = False
        else:
            titleResult = False

    returnVal = textResult or titleResult
    if keyword:
        returnVal = textResult and titleResult

    return returnVal

# 제외,포함검색어 가져오는 함수
def getKeyword(prop,kmain,idx,conn,curs):
    rKeyword = None
    with conn.cursor() as curs:
        sql = 'SELECT keyword FROM keyword_data WHERE user_idx = %s and keyword_main=%s and keyword_property = %s;'
        curs.execute(sql, (idx,kmain,prop))
        returnValue = curs.fetchall()
        rKeyword = []
        if prop == '포함':
            rKeyword.append(kmain)
        for i in range(len(returnValue)):
            rKeyword.append(str(returnValue[i]).replace("(","").replace(")","").replace(",","").replace("\'",""))
    return rKeyword

# 메인 키워드 가져오는 함수
def getMainKey(conn,key):
    result = None
    with conn.cursor() as curs:
        sql = 'SELECT keyword_main FROM `union`.keyword_data where keyword=%s;'
        curs.execute(sql, key)
        result = curs.fetchone()

    return result[0]

# 검색키워드 가져오는 함수
def getSearchKey(conn,curs):
    with conn.cursor() as curs:
        sql = 'SELECT DISTINCT keyword_main,user_idx FROM keyword_data where keyword_property=%s and user_idx in (22,23,24,25,26,27,32) and user_idx != 21;'
        curs.execute(sql, ('포함'))
        result = curs.fetchall()

    returnValue = {}
    for i in range(len(result)):
        returnValue.update({result[i][0]:{'add':[],'del':[]}})
        returnValue[result[i][0]]['add'] = getKeyword('포함',result[i][0],result[i][1],conn,curs)
        returnValue[result[i][0]]['del'] = getKeyword('제외',result[i][0],result[i][1],conn,curs)

    with conn.cursor() as curs:
        sql = "SELECT keyword_main,keyword,user_idx FROM keyword_data where keyword_property=%s and not user_idx in (22,23,24,25,26,27) and user_idx != 21 and keyword_type!=%s and keyword is not null;"
        curs.execute(sql,('포함',''))
        result = curs.fetchall()

    for i in range(len(result)):
        returnValue.update({result[i][1]:{'add':[],'del':[]}})
        returnValue[result[i][1]]['add'] = getKeyword('포함',result[i][0],result[i][2],conn,curs)
        returnValue[result[i][1]]['del'] = getKeyword('제외',result[i][0],result[i][2],conn,curs)

    return returnValue

if __name__ == '__main__':
    import pymysql
    conn = pymysql.connect(host='49.247.5.169',user='overware',password='uni1004!@',db='union',charset='utf8')
    curs = conn.cursor(pymysql.cursors.DictCursor)
    getSearchKey(conn,curs)

# 포함 키워드 가져오는 함수
def getAddKeyword(conn,curs):
    rKeyword = None
    with conn.cursor() as curs:
        sql = 'SELECT DISTINCT keyword_main,user_idx FROM keyword_data where keyword_property=%s and user_idx!=1'
        curs.execute(sql, ('포함'))
        result = curs.fetchall()
        rKeyword = []

        for i in result:
            sql = 'SELECT keyword FROM keyword_data WHERE user_idx = %s and keyword_property = %s;'
            curs.execute(sql, (i[1], '포함'))
            returnValue = curs.fetchall()

            for i in range(len(returnValue)):
                putKey = str(returnValue[i]).replace("(","").replace(")","").replace(",","").replace("\'","")
                if putKey not in rKeyword:
                    rKeyword.append(putKey)

    return rKeyword

# 기자 가져오는 함수
def getSearchReporter(conn,curs):
    returnValue = None
    with conn.cursor() as curs:
        sql = 'SELECT reporter_media_name,reporter_name FROM reporter_data;'
        curs.execute(sql)
        result = curs.fetchall()
        returnValue = []
        for i in range(len(result)):
            returnValue.append([result[i][0].replace(u'\ufeff', ''),result[i][1]])
    return returnValue

# 스타와치 키워드 가져오는 함수
def getStarWachKey(conn,curs):
    returnValue = None
    with conn.cursor() as curs:
        sql = 'SELECT distinct(keyword_main) FROM keyword_data where user_idx in (22,23,24,25,26,27);'
        curs.execute(sql)
        result = curs.fetchall()
    returnValue = [i[0] for i in result]
    return returnValue

# 메인키워드 가져오는 함수
def getMainKeyword(dicKey,title):
    returnValue = None
    for item in dicKey:
        for akey in dicKey[item]['add']:
            if title.find(akey) != -1:
                returnValue = item
    return returnValue

# 메인 뉴스키워드 가져오는 함수
def getMainNewsKeyword(conn,curs):
    result = None
    with conn.cursor() as curs:
        sql = 'SELECT k_main,k_sub,k_type FROM keyword_mail;'
        curs.execute(sql)
        result = curs.fetchall()


        returnValue = {}
        for i in range(len(result)):
            media = result[i][0].replace("\ufeff","")
            if media in returnValue:
                returnValue[media].append(result[i][1])
            else:
                returnValue.update({media:[result[i][1]]})
        # print(returnValue)

        return returnValue

# 메인 타이틀 키워드 가져오는 함수
def getMainTileKeyword(conn,curs):
    result = None
    with conn.cursor() as curs:
        sql = 'SELECT k_type_str,k_type FROM keyword_mail group by k_type_str order by k_type asc;'
        curs.execute(sql)
        result = curs.fetchall()


        returnValue = {}
        for i in range(len(result)):
            media = result[i][0].replace("\ufeff","")
            if media in returnValue:
                returnValue[media].append(result[i][1])
            else:
                returnValue.update({media:[result[i][1]]})
        # print(returnValue)

        return returnValue

#내용 체크
def checkMainNewsKeyword(content, newsKey):
    returnValue = {
        'm' : None,
        'r' : None
    }

    for s in newsKey.keys():
        if content.find(s) != -1 :
            for m in newsKey[s]:
                if content.find(m) != -1 :
                    returnValue['m'] = m
                    returnValue['r'] = s

    return returnValue

#타이틀키 체크
def checkMaintitle_key(title_key, newsKey):
    returnValue = {
        'r' : None
    }

    for s in newsKey.keys():
        if title_key.find(s) != -1 :
            returnValue['r'] = s

    return returnValue

#key_type 가져오기
def getSearchKeytpe(k_sub,k_main,conn,curs):

    with conn.cursor() as curs:
        sql = 'SELECT k_type FROM keyword_mail where k_sub = %s and k_main = %s limit 1;'
        curs.execute(sql, (k_sub,k_main))
        result = curs.fetchone()
        a = result
        if a != None:
            for b in a:
                # print(b)
                return b
        else:
            return a

#title key_type 가져오기
def getSearchTitleKeytpe(k_type_str,conn,curs):

    with conn.cursor() as curs:
        sql = 'SELECT k_type FROM keyword_mail where k_type_str = %s limit 1;'
        curs.execute(sql, (k_type_str))
        result = curs.fetchone()
        a = result
        if a != None:
            for b in a:
                # print(b)
                return b
        else:
            a == None
            return a

#Mrank 가져오기
def getSearchMrank(uid,conn,curs):

    with conn.cursor() as curs:
        sql = 'SELECT ME_rank FROM mobileent_data where uid = %s;'
        curs.execute(sql, (uid))
        result = curs.fetchone()
        a = result
        if a != None:
            for b in a:
                print(b)
                return b
        else:
            a == None
            return a

# 기자 가져오는 함수
def getReporter(conn,curs):
    result = None
    with conn.cursor() as curs:
        sql = 'SELECT reporter_media_name,reporter_name  FROM reporter_data;'
        curs.execute(sql)
        result = curs.fetchall()


        returnValue = {}
        for i in range(len(result)):
            media = result[i][0].replace("\ufeff","")
            if media in returnValue:
                returnValue[media].append(result[i][1])
            else:
                returnValue.update({media:[result[i][1]]})
        # print(returnValue)

        return returnValue

#기자 체크
def checkReporter(content, reportern):
    returnValue = {
        'm' : None,
        'r' : None
    }

    for s in reportern.keys():
        if content.find(s) != -1 :
            for m in reportern[s]:
                if content.find(m) != -1 :
                    returnValue['m'] = s
                    returnValue['r'] = m

    return returnValue

#언론사 체크
def checkMedianame(content, reportern):
    returnValue = {
        'm' : None
    }

    for s in reportern.keys():
        if content.find(s) != -1 :
            returnValue['m'] = s

    return returnValue

# 언론사, url 가져오는 함수
def getMedianameUrl(conn,curs):
    result = None
    with conn.cursor() as curs:
        sql = 'SELECT media_name, url FROM media_data where createDate >= DATE_ADD(NOW(), INTERVAL -2 day) order by createDate desc;'
        curs.execute(sql)
        result = curs.fetchall()

        returnValue = {}
        for i in range(len(result)):
            media = result[i][0].replace("\ufeff","")
            if media in returnValue:
                returnValue[media].append(result[i][1])
            else:
                returnValue.update({media:[result[i][1]]})
        # print(returnValue)

        return returnValue

#언론사, url 체크
def checkMedianameUrl(media_name, url, MedianameUrl):
    returnValue = {
        'm' : None,
        'u' : None
    }

    for s in MedianameUrl.keys():
        if media_name.find(s) != -1 :
            for m in MedianameUrl[s]:
                if url.find(m) != -1 :
                    returnValue['m'] = s
                    returnValue['u'] = m

    return returnValue

#언론사키워드 가져오는 함수
def getMediaKey(conn,curs):
    result = None
    with conn.cursor() as curs:
        sql = 'SELECT media_key, medianame  FROM medianame_data;'
        curs.execute(sql)
        result = curs.fetchall()


        returnValue = {}
        for i in range(len(result)):
            media = result[i][0].replace("\ufeff","")
            if media in returnValue:
                returnValue[media].append(result[i][1])
            else:
                returnValue.update({media:[result[i][1]]})
        # print(returnValue)

        return returnValue

#언론사키워드 가져오는 함수
def getMediaKeySub(conn,curs):
    result = None
    with conn.cursor() as curs:
        sql = 'SELECT medianame, media_key FROM medianame_data;'
        curs.execute(sql)
        result = curs.fetchall()


        returnValue = {}
        for i in range(len(result)):
            media = result[i][0].replace("\ufeff","")
            if media in returnValue:
                returnValue[media].append(result[i][1])
            else:
                returnValue.update({media:[result[i][1]]})
        # print(returnValue)

        return returnValue

#언론사키워드, url 체크
def checkMediaKeyUrl(url, MediaKey):
    returnValue = {
        'm' : None
    }

    for s in MediaKey.keys():
        if url.find(s) != -1 :
            for m in MediaKey[s]:
                returnValue['m'] = m

    return returnValue


#언론사메인 가져오는 함수
def getMediaMainname(conn,curs):
    result = None
    with conn.cursor() as curs:
        sql = 'SELECT media_name, media_subname  FROM media_main;'
        curs.execute(sql)
        result = curs.fetchall()


        returnValue = {}
        for i in range(len(result)):
            media = result[i][0].replace("\ufeff","")
            if media in returnValue:
                returnValue[media].append(result[i][1])
            else:
                returnValue.update({media:[result[i][1]]})
        # print(returnValue)

        return returnValue


#메인 미디아, 미디아 체크
def checkMediaKeyMedianame(mediaName, MediaMainname):
    check = False

    # print(MediaMainname.keys())
    for s in MediaMainname.keys():
        if mediaName.find(s) != -1 :
            check = True

    return check

# dbBackup
def dbBackup():
    dbBackup = 1

    return dbBackup
